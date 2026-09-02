import hashlib
import logging
import os
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Optional
from uuid import UUID, uuid4
from app.utils.utc import utc_now_naive

from fastapi import HTTPException
from app.utils.phone import normalize_phone_number
from app.utils.image_urls import resolve_profile_image_url
from sqlalchemy import text

from app.model.candidate_model.candidate_resume import CandidateResume
from app.config import FRONTEND_BASE_URL, JOB_ALERT_DEFAULT_TIMEZONE
from app.service import s3_service
from app.service import resume_document_service
from app.utils import resume_extraction
from app.repository.candidate_repo import CandidateProfileRepo
from app.repository.master_data_repo import MasterDataRepo
from app.service.master_data_service import MasterDataService
from app.service.subscription.subscription_validator import SubscriptionValidator
from app.service.super_admin.activity_log_service import ActivityLogService
from app.candidate_schema import (
    CandidatePersonalInfoUpdateSchema,
    CandidateSocialLinksUpdateSchema,
    SkillEntryResponse,
    CandidateSkillCreateSchema,
    CandidateSkillUpdateSchema,
    EducationEntryResponse,
    CandidateEducationCreateSchema,
    CandidateEducationUpdateSchema,
    ExperienceEntryResponse,
    CandidateExperienceCreateSchema,
    CandidateExperienceUpdateSchema,
    CandidateCertificationsUpdateSchema,
    CandidateCertificationUpdateSchema,
    CertificationEntrySchema,
    ProjectEntryResponse,
    CandidateProjectCreateSchema,
    CandidateProjectUpdateSchema,
    LanguageEntryResponse,
    CandidateLanguageCreateSchema,
    CandidateLanguageUpdateSchema,
    CandidateListingFilterParams,
    CandidateListingItemResponse,
    CandidateListingResponse,
    CandidateVisibilityUpdateSchema,
    CandidatePublicDetailResponse,
    CandidateProfileFullResponse,
    CandidateResumeDownloadResponse,
    CandidateProfessionalSnapshotUpdateSchema,
    CandidateResumeListResponse,
    CandidateResumeResponse,
    CandidateResumeUpdateSchema,
    ResumeShareLinkCreateSchema,
    ResumeShareLinkResponse,
    SharedResumeInfoResponse,
    SharedResumeAccessResponse,
    CandidateSavedJobCreateSchema,
    CandidateSavedJobListResponse,
    CandidateSavedJobResponse,
    JobAlertResponse,
    JobAlertUpsertSchema,
    JobAlertUpdateSchema,
    SavedSearchCreateSchema,
    SavedSearchResponse,
)

logger = logging.getLogger(__name__)
from sqlalchemy.ext.asyncio import AsyncSession


ALLOWED_RESUME_TYPES = {
    "application/pdf": ".pdf",
    "application/msword": ".doc",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
}
MAX_RESUME_SIZE_BYTES = 5 * 1024 * 1024

MIN_RESUME_SIZE_BYTES = 1
MAX_RESUME_FILENAME_LENGTH = 100
MAX_RESUME_COUNT = 5
JOB_ALERT_ID_PATTERN = re.compile(r"^[0-9a-fA-F]{32}$")
JOB_ALERT_TITLE_DUPLICATE_MESSAGE = (
    "An alert with this title already exists. Please use a different title."
)


def _safe_resume_extension(filename: str | None, content_type: str | None) -> str:
    expected = ALLOWED_RESUME_TYPES.get(content_type or "")
    if not expected:
        raise HTTPException(status_code=400, detail="Only PDF and Word documents are allowed")
    suffix = Path(filename or "").suffix.lower()
    if suffix and suffix != expected:
        raise HTTPException(status_code=400, detail="Resume file extension does not match the uploaded file type")
    return expected


def _validate_resume_filename_length(filename: str | None) -> None:
    # Validate only the user-facing base name (excluding the extension), so
    # the extension (.pdf/.doc/.docx) doesn't silently eat into the limit.
    stem = Path(filename or "").stem
    if len(stem) > MAX_RESUME_FILENAME_LENGTH:
        raise HTTPException(
            status_code=400,
            detail=f"File name must not exceed {MAX_RESUME_FILENAME_LENGTH} characters",
        )


_EMPTY_RESUME_DETAIL = "The selected file is empty. Please upload a resume with content."
_UNREADABLE_RESUME_DETAIL = "The uploaded file could not be read. Please upload a valid PDF or Word document."


def _docx_has_content(contents: bytes) -> bool:
    
    from docx import Document  # local import: only needed on this code path

    try:
        document = Document(BytesIO(contents))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=_UNREADABLE_RESUME_DETAIL) from exc

    if any(paragraph.text.strip() for paragraph in document.paragraphs):
        return True
    for table in document.tables:
        for row in table.rows:
            if any(cell.text.strip() for cell in row.cells):
                return True
    if document.inline_shapes and len(document.inline_shapes) > 0:
        return True
    return False


def _pdf_has_content(contents: bytes) -> bool:
    """Same idea as _docx_has_content: a PDF's byte size alone doesn't tell
    us whether any pages actually contain text or images."""
    from pypdf import PdfReader  # local import: only needed on this code path

    try:
        reader = PdfReader(BytesIO(contents))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=_UNREADABLE_RESUME_DETAIL) from exc

    if len(reader.pages) == 0:
        return False
    for page in reader.pages:
        if (page.extract_text() or "").strip():
            return True
        try:
            if len(page.images) > 0:
                return True
        except Exception:
            
            pass
    return False


def _validate_resume_has_content(contents: bytes, extension: str) -> None:
    if extension == ".docx":
        has_content = _docx_has_content(contents)
    elif extension == ".pdf":
        has_content = _pdf_has_content(contents)
    else:
        return
    if not has_content:
        raise HTTPException(status_code=400, detail=_EMPTY_RESUME_DETAIL)


def _resume_share_response(resume: CandidateResume) -> ResumeShareLinkResponse:
    share_url = None
    if resume.share_enabled and resume.share_token:
        share_url = f"{FRONTEND_BASE_URL}/cv/{resume.share_token}"
    return ResumeShareLinkResponse(
        resume_id=resume.resume_id,
        share_token=resume.share_token,
        share_enabled=resume.share_enabled,
        share_requires_email=resume.share_requires_email,
        share_expires_at=resume.share_expires_at,
        share_view_count=resume.share_view_count,
        share_url=share_url,
    )


def _calc_completion(profile, user, resume_detail) -> int:
    score = 0
    if user.first_name and user.last_name:
        score += 10
    if user.mobile_number:
        score += 10
    if profile.headline:
        score += 10
    if profile.summary:
        score += 10
    if profile.current_location:
        score += 5
    if resume_detail and resume_detail.skills_json and resume_detail.skills_json.get("skills"):
        score += 5
    if profile.linkedin_url or profile.github_url or profile.dribbble_url or profile.twitter_url:
        score += 10
    if resume_detail:
        if resume_detail.education_json and resume_detail.education_json.get("education"):
            score += 15
        if resume_detail.experience_json and resume_detail.experience_json.get("experience"):
            score += 15
        if resume_detail.certifications_json and resume_detail.certifications_json.get("certifications"):
            score += 10
    return min(score, 100)


def _display_image_url(stored_value: str | None) -> str | None:
    return resolve_profile_image_url(stored_value)


def _resume_list_section(resume_detail, attr: str, key: str) -> list[dict] | None:
    value = getattr(resume_detail, attr, None) if resume_detail else None
    if isinstance(value, list):
        items = [item for item in value if isinstance(item, dict)]
        return items or None
    if not isinstance(value, dict):
        return None
    section = value.get(key)
    if not isinstance(section, list):
        return None
    return [item for item in section if isinstance(item, dict)]


def _raw_resume_list_section(resume_detail, attr: str, key: str) -> list:
    value = getattr(resume_detail, attr, None) if resume_detail else None
    if isinstance(value, list):
        return value
    if isinstance(value, dict):
        section = value.get(key)
        return section if isinstance(section, list) else []
    return []


def _legacy_entry_id(section: str, index: int, item: Any) -> str:
    digest = hashlib.sha1(repr(item).encode("utf-8")).hexdigest()[:16]
    return f"legacy-{section}-{index}-{digest}"


def _profile_section_entries(
    raw: list,
    *,
    id_field: str,
    section: str,
    required_fields: tuple[str, ...],
    string_field: str | None = None,
    defaults: dict | None = None,
) -> list[dict]:
    items = []
    for index, value in enumerate(raw):
        if isinstance(value, str) and string_field:
            value = {string_field: value}
        if not isinstance(value, dict):
            continue
        item = _clean_dict({**(defaults or {}), **dict(value)})
        if any(not item.get(field) for field in required_fields):
            continue
        item[id_field] = item.get(id_field) or _legacy_entry_id(section, index, item)
        items.append(item)
    return items


def _resume_skills_section(resume_detail) -> dict | None:
    value = getattr(resume_detail, "skills_json", None) if resume_detail else None
    if isinstance(value, dict):
        skills = value.get("skills")
        if isinstance(skills, list):
            return {"skills": skills}
        return value
    if isinstance(value, list):
        return {"skills": value}
    return None


def _none_if_blank(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _bool_or_default(value: Any, default: bool = False) -> bool:
    return default if value is None else bool(value)


def _parse_date(value: Any) -> date | None:
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    if not value:
        return None
    text_value = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m", "%Y"):
        try:
            parsed = datetime.strptime(text_value, fmt)
            return parsed.date()
        except ValueError:
            continue
    return None


def _date_sort_value(value: Any) -> date:
    return _parse_date(value) or date.min


def _duration_between(start_value: Any, end_value: Any, currently_working: bool = False) -> str | None:
    start = _parse_date(start_value)
    if not start:
        return None
    end = date.today() if currently_working else (_parse_date(end_value) or date.today())
    if end < start:
        return None
    months = (end.year - start.year) * 12 + end.month - start.month
    years, remaining_months = divmod(max(months, 0), 12)
    parts = []
    if years:
        parts.append(f"{years} year{'s' if years != 1 else ''}")
    if remaining_months:
        parts.append(f"{remaining_months} month{'s' if remaining_months != 1 else ''}")
    return ", ".join(parts) or "Less than 1 month"


def _clean_dict(item: dict) -> dict:
    return {key: _none_if_blank(value) for key, value in item.items()}


def _normalized_list_section(resume_detail, attr: str, key: str) -> list[dict] | None:
    items = _resume_list_section(resume_detail, attr, key)
    if items is None:
        return None
    return [_clean_dict(item) for item in items if isinstance(item, dict)]


def _normalize_experience_items(resume_detail) -> list[dict] | None:
    items = _normalized_list_section(resume_detail, "experience_json", "experience")
    if items is None:
        return None
    normalized = []
    for item in items:
        currently_working = bool(
            item.get("currently_working")
            or item.get("current_company_flag")
            or item.get("is_current")
        )
        start_date = _none_if_blank(item.get("start_date"))
        end_date = None if currently_working else _none_if_blank(item.get("end_date"))
        normalized.append(
            {
                **item,
                "company": item.get("company"),
                "designation": item.get("designation") or item.get("role") or item.get("title"),
                "role": item.get("role") or item.get("designation") or item.get("title"),
                "employment_type": item.get("employment_type"),
                "location": item.get("location"),
                "start_date": start_date,
                "end_date": end_date,
                "current_company_flag": currently_working,
                "duration": _duration_between(start_date, end_date, currently_working),
                "responsibilities": item.get("responsibilities") or item.get("key_highlights"),
                "achievements": item.get("achievements"),
            }
        )
    return sorted(
        normalized,
        key=lambda item: (
            bool(item.get("current_company_flag")),
            _date_sort_value(item.get("end_date") or item.get("start_date")),
        ),
        reverse=True,
    )


def _normalize_education_items(resume_detail) -> list[dict] | None:
    items = _normalized_list_section(resume_detail, "education_json", "education")
    if items is None:
        return None
    return sorted(
        items,
        key=lambda item: _date_sort_value(
            item.get("graduation_year")
            or item.get("end_year")
            or item.get("end_date")
            or item.get("start_year")
        ),
        reverse=True,
    )


def _normalize_named_items(resume_detail, attr: str, key: str) -> list[dict] | None:
    items = _normalized_list_section(resume_detail, attr, key)
    if items is None:
        return None
    return sorted(
        items,
        key=lambda item: str(
            item.get("issue_date")
            or item.get("start_date")
            or item.get("name")
            or item.get("title")
            or ""
        ),
        reverse=True,
    )


def _normalize_skills(profile, resume_detail) -> dict:
    raw = []
    if resume_detail and isinstance(resume_detail.skills_json, dict):
        raw = resume_detail.skills_json.get("skills") or []
    elif resume_detail and isinstance(resume_detail.skills_json, list):
        raw = resume_detail.skills_json
    if not raw and getattr(profile, "skills_summary", None):
        raw = [part.strip() for part in str(profile.skills_summary).split(",")]

    seen = set()
    skills = []
    categories: dict[str, list[dict]] = {}
    for value in raw:
        item = value if isinstance(value, dict) else {"name": value}
        name = _none_if_blank(item.get("name") or item.get("skill") or item.get("title"))
        if not name:
            continue
        key = str(name).casefold()
        if key in seen:
            continue
        seen.add(key)
        cleaned = _clean_dict({**item, "name": name})
        response_item = cleaned if isinstance(value, dict) else name
        skills.append(response_item)
        category = _none_if_blank(cleaned.get("category"))
        if category:
            categories.setdefault(str(category), []).append(cleaned)
    response = {"skills": skills}
    if categories:
        response["categories"] = categories
    return response


def _current_employment(profile, experience: list[dict] | None) -> dict | None:
    experience = experience or []
    current = next((item for item in experience if item.get("current_company_flag")), None)
    current = current or (experience[0] if experience else None)
    if not current and not getattr(profile, "current_company", None):
        return None
    return {
        "company": _none_if_blank(getattr(profile, "current_company", None))
        or (current or {}).get("company"),
        "designation": (current or {}).get("designation") or (current or {}).get("role"),
        "employment_type": (current or {}).get("employment_type")
        or _none_if_blank(getattr(profile, "desired_employment", None)),
        "location": (current or {}).get("location")
        or _none_if_blank(getattr(profile, "current_location", None)),
        "start_date": (current or {}).get("start_date"),
        "duration": (current or {}).get("duration"),
    }


def _split_locations(value: Any) -> list[str] | None:
    value = _none_if_blank(value)
    if not value:
        return None
    if isinstance(value, list):
        return [item for item in (_none_if_blank(item) for item in value) if item]
    return [item.strip() for item in str(value).split(",") if item.strip()] or None


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _resume_file_type(resume) -> str | None:
    name = getattr(resume, "file_name", None)
    suffix = Path(name or "").suffix.lower().lstrip(".")
    return suffix or None


def _resume_metadata(resume, application_id: str | None = None) -> dict | None:
    if not resume:
        return None
    preview_url = None
    download_url = None
    if application_id:
        preview_url = f"/employer/job-applicants/{application_id}/resume/preview-file"
        download_url = f"/employer/job-applicants/{application_id}/resume/download-file"
    return {
        "resume_id": getattr(resume, "resume_id", None),
        "file_name": _none_if_blank(getattr(resume, "file_name", None)),
        "filename": _none_if_blank(getattr(resume, "file_name", None)),
        "uploaded_at": getattr(resume, "uploaded_at", None),
        "upload_date": getattr(resume, "uploaded_at", None),
        "download_url": download_url,
        "preview_url": preview_url,
        "file_size": getattr(resume, "file_size", None),
        "file_type": _resume_file_type(resume),
        "is_active": getattr(resume, "is_active", False),
        "version_name": _none_if_blank(getattr(resume, "version_name", None)),
        "template": _none_if_blank(getattr(resume, "template", None)),
    }


def _profile_completion(profile, user, resume, sections: dict[str, Any]) -> dict:
    checks = {
        "full_name": bool(CandidateProfileService._full_name(user)),
        "email": bool(_none_if_blank(getattr(user, "email", None))),
        "phone": bool(_none_if_blank(getattr(user, "mobile_number", None))),
        "headline": bool(_none_if_blank(getattr(profile, "headline", None))),
        "summary": bool(_none_if_blank(getattr(profile, "summary", None))),
        "location": bool(_none_if_blank(getattr(profile, "current_location", None))),
        "current_company": bool(_none_if_blank(getattr(profile, "current_company", None))),
        "experience": bool(sections.get("experience")),
        "education": bool(sections.get("education")),
        "skills": bool((sections.get("skills") or {}).get("skills")),
        "resume": bool(resume),
    }
    completed = [key for key, done in checks.items() if done]
    missing = [key for key, done in checks.items() if not done]
    percentage = round((len(completed) / len(checks)) * 100) if checks else 0
    return {
        "percentage": percentage,
        "missing_fields": missing,
        "completed_sections": completed,
    }


async def _get_profile_or_404(session: AsyncSession, user_id: UUID):
    profile = await CandidateProfileRepo.get_profile_by_user_id(session, user_id)
    if profile:
        return profile

    user = await CandidateProfileRepo.get_user_by_id(session, user_id)
    roles = getattr(user, "roles", []) if user else []
    is_candidate = any(
        getattr(role, "role_code", None) == "ROLE_CANDIDATE"
        for role in roles
    )

    if user and is_candidate:
        return await CandidateProfileRepo.create_profile_for_user(session, user_id)

    raise HTTPException(status_code=404, detail="Candidate profile not found")


async def _get_user_or_404(session: AsyncSession, user_id: UUID):
    user = await CandidateProfileRepo.get_user_by_id(session, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user


async def _require_candidate_feature(
    session: AsyncSession,
    user_id: UUID,
    feature_name: str,
) -> None:
    await SubscriptionValidator(
        session=session,
        user_id=user_id,
        role="CANDIDATE",
    ).require_feature(feature_name)


class CandidateProfileService:

    @staticmethod
    def _full_name(user) -> str:
        return " ".join(
            part
            for part in (
                getattr(user, "first_name", None),
                getattr(user, "middle_name", None),
                getattr(user, "last_name", None),
            )
            if part
        )

    @staticmethod
    def _candidate_listing_item(profile, user) -> CandidateListingItemResponse:
        return CandidateListingItemResponse(
            candidate_id=profile.candidate_id,
            full_name=CandidateProfileService._full_name(user),
            headline=profile.headline,
            current_company=profile.current_company,
            current_location=profile.current_location,
            total_experience=profile.total_experience,
            experience_level=profile.experience_level,
            skills_summary=profile.skills_summary,
            work_preference=profile.work_preference,
            open_to_work=_bool_or_default(profile.open_to_work, False),
            profile_completion_pct=int(profile.profile_completion_pct or 0),
            active_resume_id=profile.active_resume_id,
            profile_image_url=_display_image_url(getattr(user, "profile_image_url", None)),
        )

    @staticmethod
    async def list_public_candidates(session: AsyncSession, filters: CandidateListingFilterParams) -> CandidateListingResponse:
        total, rows = await CandidateProfileRepo.list_public_candidates(
            session=session,
            search=filters.search,
            location=filters.location,
            skill=filters.skill,
            experience_level=filters.experience_level,
            work_preference=filters.work_preference,
            page=filters.page,
            page_size=filters.page_size,
        )
        return CandidateListingResponse(
            total=total,
            page=filters.page,
            page_size=filters.page_size,
            items=[
                CandidateProfileService._candidate_listing_item(profile, user)
                for profile, user in rows
            ],
        )

    @staticmethod
    async def get_public_candidate_detail(session: AsyncSession, candidate_id: str) -> CandidatePublicDetailResponse:
        row = await CandidateProfileRepo.get_public_candidate_detail(session, candidate_id)
        if not row:
            raise HTTPException(status_code=404, detail="Candidate not found")
        profile, user = row
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        item = CandidateProfileService._candidate_listing_item(profile, user)
        return CandidatePublicDetailResponse(
            **item.model_dump(),
            professional_title=profile.headline,
            about_you=profile.summary,
            location=profile.current_location,
            cover_image_url=_display_image_url(getattr(user, "cover_image_url", None)),
            profile_visibility=profile.profile_visibility or "PRIVATE",
            searchable_flag=_bool_or_default(profile.searchable_flag, True),
            search_engine_indexing=_bool_or_default(getattr(profile, "search_engine_indexing", None), False),
            summary=profile.summary,
            preferred_location=profile.preferred_location,
            desired_employment=profile.desired_employment,
            salary_expectation=profile.salary_expectation,
            target_roles=profile.target_roles,
            education=_resume_list_section(resume_detail, "education_json", "education"),
            experience=_resume_list_section(resume_detail, "experience_json", "experience"),
            skills=_resume_skills_section(resume_detail),
            certifications=_resume_list_section(resume_detail, "certifications_json", "certifications"),
            projects=_resume_list_section(resume_detail, "projects_json", "projects"),
            languages=_resume_list_section(resume_detail, "languages_json", "languages"),
            resumes=(
                [{"resume_id": profile.active_resume_id, "is_active": True}]
                if profile.active_resume_id
                else []
            ),
        )

    @staticmethod
    async def get_candidate_detail_for_employer(
        session: AsyncSession,
        candidate_id: str,
        employer_user_id: Optional[UUID] = None,
    ) -> CandidatePublicDetailResponse:
        row = await CandidateProfileRepo.get_candidate_detail_for_employer(session, candidate_id)
        if not row:
            profile = await CandidateProfileRepo.get_candidate_profile_by_id(
                session=session,
                candidate_id=candidate_id,
            )
            if profile:
                raise HTTPException(
                    status_code=403,
                    detail="Candidate public profile is private or not available.",
                )
            raise HTTPException(status_code=404, detail="Candidate not found.")
        profile, user = row
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        employer_context = (
            await CandidateProfileRepo.get_employer_candidate_context(
                session=session,
                candidate_id=profile.candidate_id,
                employer_user_id=employer_user_id,
            )
            if employer_user_id
            else {
                "saved_candidate": False,
                "application": None,
                "interview": None,
                "notes_count": 0,
            }
        )
        application = employer_context.get("application")
        interview = employer_context.get("interview")
        resume = await CandidateProfileRepo.get_default_resume_for_employer_profile(
            session=session,
            candidate_id=profile.candidate_id,
            active_resume_id=profile.active_resume_id,
        )
        experience = _normalize_experience_items(resume_detail)
        education = _normalize_education_items(resume_detail)
        skills = _normalize_skills(profile, resume_detail)
        certifications = _normalize_named_items(
            resume_detail,
            "certifications_json",
            "certifications",
        )
        projects = _normalize_named_items(resume_detail, "projects_json", "projects")
        languages = _normalize_named_items(resume_detail, "languages_json", "languages")
        current_employment = _current_employment(profile, experience)
        current_designation = (
            (current_employment or {}).get("designation")
            if current_employment
            else None
        )
        resume_section = _resume_metadata(
            resume,
            getattr(application, "application_id", None),
        )
        completion = _profile_completion(
            profile,
            user,
            resume,
            {"experience": experience, "education": education, "skills": skills},
        )
        item = CandidateProfileService._candidate_listing_item(profile, user)
        profile_photo = _display_image_url(getattr(user, "profile_image_url", None))
        location = _none_if_blank(profile.current_location)
        base_item = item.model_dump(
            exclude={
                "profile_image_url",
                "headline",
                "current_company",
                "work_preference",
                "open_to_work",
                "profile_completion_pct",
            }
        )
        return CandidatePublicDetailResponse(
            **base_item,
            profile_image_url=profile_photo,
            profile_photo=profile_photo,
            headline=_none_if_blank(profile.headline),
            email=getattr(user, "email", None),
            phone_number=getattr(user, "mobile_number", None),
            phone=getattr(user, "mobile_number", None),
            professional_title=_none_if_blank(getattr(profile, "headline", None)),
            professional_summary=_none_if_blank(getattr(profile, "summary", None)),
            about_you=_none_if_blank(getattr(profile, "summary", None)),
            about_me=_none_if_blank(getattr(profile, "summary", None)),
            location=location,
            city=location,
            cover_image_url=_display_image_url(getattr(user, "cover_image_url", None)),
            profile_visibility=getattr(profile, "profile_visibility", None) or "PRIVATE",
            searchable_flag=_bool_or_default(getattr(profile, "searchable_flag", None), True),
            search_engine_indexing=_bool_or_default(getattr(profile, "search_engine_indexing", None), False),
            summary=_none_if_blank(getattr(profile, "summary", None)),
            current_company=_none_if_blank(getattr(profile, "current_company", None))
            or (current_employment or {}).get("company"),
            current_designation=_none_if_blank(current_designation),
            total_experience_years=_to_float(getattr(profile, "total_experience", None)),
            current_employment=current_employment,
            notice_period=_none_if_blank(getattr(profile, "notice_period", None)),
            employment_type=_none_if_blank(getattr(profile, "desired_employment", None)),
            work_preference=_none_if_blank(getattr(profile, "work_preference", None)),
            preferred_locations=_split_locations(getattr(profile, "preferred_location", None)),
            current_salary=_to_float(getattr(profile, "current_ctc", None)),
            expected_salary=_to_float(getattr(profile, "expected_ctc", None)),
            open_to_work=_bool_or_default(getattr(profile, "open_to_work", None), False),
            remote_preference=(
                getattr(profile, "work_preference", None)
                if str(getattr(profile, "work_preference", None) or "").upper() == "REMOTE"
                else None
            ),
            relocation_preference=None,
            preferred_location=_none_if_blank(getattr(profile, "preferred_location", None)),
            desired_employment=_none_if_blank(getattr(profile, "desired_employment", None)),
            salary_expectation=_none_if_blank(getattr(profile, "salary_expectation", None)),
            target_roles=_none_if_blank(getattr(profile, "target_roles", None)),
            education=education,
            experience=experience,
            skills=skills,
            certifications=certifications,
            projects=projects,
            languages=languages,
            resume=resume_section,
            resumes=[resume_section] if resume_section else [],
            linkedin_url=_none_if_blank(getattr(profile, "linkedin_url", None)),
            github_url=_none_if_blank(getattr(profile, "github_url", None)),
            portfolio_url=_none_if_blank(getattr(profile, "portfolio_url", None)),
            website=_none_if_blank(getattr(profile, "website_url", None)),
            leetcode_url=None,
            hackerrank_url=None,
            social_links={
                "linkedin": _none_if_blank(getattr(profile, "linkedin_url", None)),
                "github": _none_if_blank(getattr(profile, "github_url", None)),
                "portfolio": _none_if_blank(getattr(profile, "portfolio_url", None)),
                "website": _none_if_blank(getattr(profile, "website_url", None)),
                "leetcode": None,
                "hackerrank": None,
            },
            profile_completion=completion,
            profile_completion_pct=completion["percentage"],
            saved_candidate=bool(employer_context.get("saved_candidate")),
            applied_job=(
                {
                    "application_id": application.application_id,
                    "job_id": application.job_id,
                    "applied_at": application.applied_at,
                    "status": application.application_status,
                }
                if application
                else None
            ),
            current_application_status=(
                application.application_status if application else None
            ),
            interview_status=interview.status if interview else None,
            interview_date=interview.interview_date if interview else None,
            notes_count=int(employer_context.get("notes_count") or 0),
        )

    @staticmethod
    async def get_full_profile(session: AsyncSession, user_id: UUID) -> CandidateProfileFullResponse:
        user = await _get_user_or_404(session, user_id)
        profile = await _get_profile_or_404(session, user_id)
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)

        education = []
        experience = []
        skills = []
        certifications = None
        projects = []
        languages = []

        if resume_detail:
            education = _profile_section_entries(
                _raw_resume_list_section(resume_detail, "education_json", "education"),
                id_field="education_id",
                section="education",
                required_fields=("institution",),
            )
            experience = _profile_section_entries(
                _raw_resume_list_section(resume_detail, "experience_json", "experience"),
                id_field="experience_id",
                section="experience",
                required_fields=("company", "role"),
            )
            skills = _profile_section_entries(
                _raw_resume_list_section(resume_detail, "skills_json", "skills"),
                id_field="skill_id",
                section="skill",
                required_fields=("name",),
                string_field="name",
            )
            certifications = _raw_resume_list_section(
                resume_detail,
                "certifications_json",
                "certifications",
            )
            certifications = [
                item for item in certifications if isinstance(item, dict)
            ] or None
            projects = _profile_section_entries(
                _raw_resume_list_section(resume_detail, "projects_json", "projects"),
                id_field="project_id",
                section="project",
                required_fields=("title",),
            )
            languages = _profile_section_entries(
                _raw_resume_list_section(resume_detail, "languages_json", "languages"),
                id_field="language_id",
                section="language",
                required_fields=("name",),
                string_field="name",
                defaults={"proficiency_level": "BASIC"},
            )

        country = await MasterDataRepo.get_country(session, profile.country_id) if profile.country_id else None
        primary_location = await MasterDataRepo.get_location(session, profile.primary_location_id) if profile.primary_location_id else None
        preferred_location = await MasterDataRepo.get_location(session, profile.preferred_location_id) if profile.preferred_location_id else None
        notice_period = await MasterDataRepo.get_notice_period(session, profile.notice_period_id) if profile.notice_period_id else None
        salary_expectation = await MasterDataRepo.get_salary_expectation(session, profile.salary_expectation_id) if profile.salary_expectation_id else None
        target_role_rows = await MasterDataRepo.get_candidate_target_roles(session, profile.candidate_id)

        return CandidateProfileFullResponse(
            user_id=user.user_id,
            first_name=user.first_name,
            middle_name=user.middle_name,
            last_name=user.last_name,
            email=user.email,
            phone_number=user.mobile_number,
            phone_country_code=getattr(user, "country_code", None),
            phone_country_iso2=getattr(user, "phone_country_iso2", None),
            profile_image_url=_display_image_url(user.profile_image_url),
            cover_image_url=_display_image_url(user.cover_image_url),
            candidate_id=profile.candidate_id,
            professional_title=profile.headline,
            about_you=profile.summary,
            country_id=profile.country_id,
            country=country.name if country else None,
            primary_location_id=profile.primary_location_id,
            location_id=profile.primary_location_id,
            primary_location=primary_location.name if primary_location else profile.current_location,
            preferred_location_id=profile.preferred_location_id,
            preferred_locations=preferred_location.name if preferred_location else profile.preferred_location,
            website=profile.website_url,
            portfolio_url=profile.portfolio_url,
            experience_level=profile.experience_level,
            current_company=profile.current_company,
            notice_period_id=profile.notice_period_id,
            notice_period=notice_period.label if notice_period else profile.notice_period,
            desired_employment=profile.desired_employment,
            salary_expectation_id=profile.salary_expectation_id,
            salary_expectation=salary_expectation.label if salary_expectation else profile.salary_expectation,
            current_ctc=profile.current_ctc,
            expected_ctc=profile.expected_ctc,
            work_preference=profile.work_preference,
            target_role_ids=[r.target_role_id for r in target_role_rows],
            target_roles=[r.name for r in target_role_rows],
            linkedin_url=profile.linkedin_url,
            dribbble_url=profile.dribbble_url,
            github_url=profile.github_url,
            twitter_url=profile.twitter_url,
            profile_visibility=profile.profile_visibility or "PRIVATE",
            searchable_flag=_bool_or_default(profile.searchable_flag, True),
            open_to_work=_bool_or_default(profile.open_to_work, False),
            search_engine_indexing=_bool_or_default(getattr(profile, "search_engine_indexing", None), False),
            profile_completion_pct=int(profile.profile_completion_pct or 0),
            active_resume_id=profile.active_resume_id,
            education=education,
            experience=experience,
            skills=skills,
            certifications=certifications,
            projects=projects,
            languages=languages,
        )

    @staticmethod
    async def update_personal_info(session: AsyncSession, user_id: UUID, data: CandidatePersonalInfoUpdateSchema) -> dict:
        user = await _get_user_or_404(session, user_id)
        profile = await _get_profile_or_404(session, user_id)

        # Build kwargs only for fields explicitly sent in the request (partial update)
        sent = data.model_fields_set
        user_kwargs = {}
        if "first_name" in sent and data.first_name is not None:
            user_kwargs["first_name"] = data.first_name
        if "last_name" in sent and data.last_name is not None:
            user_kwargs["last_name"] = data.last_name
        if "middle_name" in sent:
            user_kwargs["middle_name"] = data.middle_name
        if "phone_number" in sent and data.phone_number is not None:
            # (so just editing the phone number without touching the country
            # dropdown keeps validating against their existing country).
            effective_country_code = (
                data.phone_country_code if "phone_country_code" in sent and data.phone_country_code
                else (getattr(user, "country_code", None) or "+91")
            )
            effective_country_iso2 = (
                data.phone_country_iso2 if "phone_country_iso2" in sent and data.phone_country_iso2
                else getattr(user, "phone_country_iso2", None)
            )
            try:
                normalized = normalize_phone_number(effective_country_code, data.phone_number, effective_country_iso2)
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=str(exc))
            user_kwargs["mobile_number"] = normalized.national_number
            user_kwargs["country_code"] = normalized.country_code
            if "phone_country_iso2" in sent:
                user_kwargs["phone_country_iso2"] = data.phone_country_iso2
        elif "phone_country_code" in sent and data.phone_country_code is not None:
            # Country changed but phone number wasn't resent -- validate the
            # existing saved number against the new country.
            effective_country_iso2 = (
                data.phone_country_iso2 if "phone_country_iso2" in sent and data.phone_country_iso2
                else getattr(user, "phone_country_iso2", None)
            )
            try:
                normalized = normalize_phone_number(data.phone_country_code, user.mobile_number, effective_country_iso2)
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=str(exc))
            user_kwargs["country_code"] = normalized.country_code
            if "phone_country_iso2" in sent:
                user_kwargs["phone_country_iso2"] = data.phone_country_iso2
        elif "phone_country_iso2" in sent and data.phone_country_iso2 is not None:
            effective_country_code = getattr(user, "country_code", None) or "+91"
            try:
                normalize_phone_number(effective_country_code, user.mobile_number, data.phone_country_iso2)
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=str(exc))
            user_kwargs["phone_country_iso2"] = data.phone_country_iso2

        if user_kwargs:
            await CandidateProfileRepo.update_user(
                session=session,
                user_id=user_id,
                updated_by=str(user_id),
                **user_kwargs,
            )
            for attr, val in user_kwargs.items():
                setattr(user, attr, val)

        # Map frontend field names -> underlying CandidateProfile column names
        field_to_column = {
            "professional_title": "headline",
            "about_you": "summary",
            "primary_location": "current_location",
            "preferred_locations": "preferred_location",
            "website": "website_url",
            "portfolio_url": "portfolio_url",
        }
        profile_kwargs = {}
        for field, column in field_to_column.items():
            if field in sent:
                profile_kwargs[column] = getattr(data, field, None)

        
        if "preferred_locations" in sent:
            profile_kwargs["preferred_location_id"] = None

        # Country dropdown drives which Location options are valid.
        effective_country_id = profile.country_id
        if "country_id" in sent and data.country_id is not None:
            country = await MasterDataService.validate_country(session, data.country_id)
            profile_kwargs["country_id"] = country.country_id
            effective_country_id = country.country_id


        requested_primary_location_id = None
        primary_location_explicitly_cleared = False
        if "location_id" in sent:
            if data.location_id is not None:
                requested_primary_location_id = data.location_id
            else:
                primary_location_explicitly_cleared = True
        if "primary_location_id" in sent:
            if data.primary_location_id is not None:
                if requested_primary_location_id and requested_primary_location_id != data.primary_location_id:
                    raise HTTPException(status_code=400, detail="location_id and primary_location_id must match when both are provided")
                requested_primary_location_id = data.primary_location_id
                primary_location_explicitly_cleared = False
            else:
                primary_location_explicitly_cleared = True

        if requested_primary_location_id:
            location = await MasterDataService.validate_location_for_country(
                session, requested_primary_location_id, effective_country_id
            )
            profile_kwargs["primary_location_id"] = location.location_id
            profile_kwargs["current_location"] = location.name  # keep display-name cache in sync
            if not effective_country_id:
                profile_kwargs["country_id"] = location.country_id
                effective_country_id = location.country_id
        elif primary_location_explicitly_cleared:
            profile_kwargs["primary_location_id"] = None
            profile_kwargs["current_location"] = None

        if "preferred_location_id" in sent and data.preferred_location_id is not None:
            location = await MasterDataService.validate_location_for_country(
                session, data.preferred_location_id, effective_country_id
            )
            profile_kwargs["preferred_location_id"] = location.location_id
            profile_kwargs["preferred_location"] = location.name  # keep display-name cache in sync
            if not effective_country_id:
                profile_kwargs["country_id"] = location.country_id
                effective_country_id = location.country_id

        for k, v in profile_kwargs.items():
            setattr(profile, k, v)

        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        profile_kwargs["profile_completion_pct"] = completion

        await CandidateProfileRepo.update_profile(
            session=session,
            candidate_id=profile.candidate_id,
            updated_by=str(user_id),
            **profile_kwargs,
        )
        await ActivityLogService.create_log_for_user_id(
            session=session,
            user_id=user_id,
            actor_role="CANDIDATE",
            action="PROFILE_UPDATED",
            entity_type="CandidateProfile",
            entity_id=profile.candidate_id,
            description="Candidate personal info updated",
        )

        return {"message": "Personal info updated successfully", "profile_completion_pct": completion}

    @staticmethod
    async def update_social_links(session: AsyncSession, user_id: UUID, data: CandidateSocialLinksUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        sent = data.model_fields_set
        update_kwargs = {field: getattr(data, field) for field in sent}
        if not update_kwargs:
            return {"message": "No changes provided"}

        user = await _get_user_or_404(session, user_id)
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        for k, v in update_kwargs.items():
            setattr(profile, k, v)
        completion = _calc_completion(profile, user, resume_detail)
        update_kwargs["profile_completion_pct"] = completion

        await CandidateProfileRepo.update_profile(
            session=session,
            candidate_id=profile.candidate_id,
            updated_by=str(user_id),
            **update_kwargs,
        )
        await ActivityLogService.create_log_for_user_id(
            session=session,
            user_id=user_id,
            actor_role="CANDIDATE",
            action="PROFILE_UPDATED",
            entity_type="CandidateProfile",
            entity_id=profile.candidate_id,
            description="Candidate social links updated",
        )
        return {"message": "Social links updated successfully", "profile_completion_pct": completion}

    # ── Skills (Add Skill modal: per-entry CRUD) ──────────────────────────

    @staticmethod
    async def add_skill(session: AsyncSession, user_id: UUID, data: CandidateSkillCreateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)

        existing_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        existing_skills = (existing_detail.skills_json or {}).get("skills", []) if existing_detail else []
        new_name = (data.name or "").strip().lower()
        if any((entry.get("name") or "").strip().lower() == new_name for entry in existing_skills):
            raise HTTPException(status_code=400, detail="This skill has already been added.")

        item = data.model_dump()
        item["skill_id"] = uuid4().hex
        await CandidateProfileRepo.add_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="skills_json",
            list_key="skills",
            item=item,
        )
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Skill added successfully", "profile_completion_pct": completion, "skill": item}

    @staticmethod
    async def update_skill(session: AsyncSession, user_id: UUID, skill_id: str, data: CandidateSkillUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)

        if data.name is not None:
            existing_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
            existing_skills = (existing_detail.skills_json or {}).get("skills", []) if existing_detail else []
            new_name = data.name.strip().lower()
            if any(
                (entry.get("name") or "").strip().lower() == new_name and entry.get("skill_id") != skill_id
                for entry in existing_skills
            ):
                raise HTTPException(status_code=400, detail="This skill has already been added.")

        updated = await CandidateProfileRepo.update_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="skills_json",
            list_key="skills",
            item_id_field="skill_id",
            item_id=skill_id,
            updates=data.model_dump(exclude_unset=True),
        )
        if updated is None:
            raise HTTPException(status_code=404, detail="Skill not found")
        return {"message": "Skill updated successfully", "skill": updated}

    @staticmethod
    async def delete_skill(session: AsyncSession, user_id: UUID, skill_id: str) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)
        deleted = await CandidateProfileRepo.delete_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="skills_json",
            list_key="skills",
            item_id_field="skill_id",
            item_id=skill_id,
        )
        if not deleted:
            raise HTTPException(status_code=404, detail="Skill not found")
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Skill removed successfully", "profile_completion_pct": completion}

    # ── Education (Add Education modal: per-entry CRUD) ───────────────────

    @staticmethod
    async def add_education(session: AsyncSession, user_id: UUID, data: CandidateEducationCreateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)

        item = data.model_dump()
        item["education_id"] = uuid4().hex
        await CandidateProfileRepo.add_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="education_json",
            list_key="education",
            item=item,
        )
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Education added successfully", "profile_completion_pct": completion, "education": item}

    @staticmethod
    async def update_education(session: AsyncSession, user_id: UUID, education_id: str, data: CandidateEducationUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        updated = await CandidateProfileRepo.update_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="education_json",
            list_key="education",
            item_id_field="education_id",
            item_id=education_id,
            updates=data.model_dump(exclude_unset=True),
        )
        if updated is None:
            raise HTTPException(status_code=404, detail="Education entry not found")
        return {"message": "Education updated successfully", "education": updated}

    @staticmethod
    async def delete_education(session: AsyncSession, user_id: UUID, education_id: str) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)
        deleted = await CandidateProfileRepo.delete_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="education_json",
            list_key="education",
            item_id_field="education_id",
            item_id=education_id,
        )
        if not deleted:
            raise HTTPException(status_code=404, detail="Education entry not found")
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Education removed successfully", "profile_completion_pct": completion}

    # ── Experience (Add Experience modal: per-entry CRUD) ──────────────────

    @staticmethod
    async def add_experience(session: AsyncSession, user_id: UUID, data: CandidateExperienceCreateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)

        item = data.model_dump(mode="json")
        item["experience_id"] = uuid4().hex
        await CandidateProfileRepo.add_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="experience_json",
            list_key="experience",
            item=item,
        )
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Experience added successfully", "profile_completion_pct": completion, "experience": item}

    @staticmethod
    async def update_experience(session: AsyncSession, user_id: UUID, experience_id: str, data: CandidateExperienceUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        updated = await CandidateProfileRepo.update_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="experience_json",
            list_key="experience",
            item_id_field="experience_id",
            item_id=experience_id,
            updates=data.model_dump(mode="json", exclude_unset=True),
        )
        if updated is None:
            raise HTTPException(status_code=404, detail="Experience entry not found")
        return {"message": "Experience updated successfully", "experience": updated}

    @staticmethod
    async def delete_experience(session: AsyncSession, user_id: UUID, experience_id: str) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)
        deleted = await CandidateProfileRepo.delete_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="experience_json",
            list_key="experience",
            item_id_field="experience_id",
            item_id=experience_id,
        )
        if not deleted:
            raise HTTPException(status_code=404, detail="Experience entry not found")
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Experience removed successfully", "profile_completion_pct": completion}

    @staticmethod
    async def update_certifications(session: AsyncSession, user_id: UUID, data: CandidateCertificationsUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)

        certifications_json = {"certifications": [c.model_dump() for c in data.certifications]}

        resume_detail = await CandidateProfileRepo.upsert_resume_detail(
            session=session,
            candidate_id=profile.candidate_id,
            certifications_json=certifications_json,
        )
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session,
            candidate_id=profile.candidate_id,
            updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {
            "message": "Certifications updated successfully",
            "profile_completion_pct": completion,
            "certifications": certifications_json["certifications"],
        }

    @staticmethod
    async def add_certification(session: AsyncSession, user_id: UUID, data: CertificationEntrySchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)

        item = data.model_dump()
        item["certification_id"] = uuid4().hex
        await CandidateProfileRepo.add_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="certifications_json",
            list_key="certifications",
            item=item,
        )
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Certification added successfully", "profile_completion_pct": completion, "certification": item}

    @staticmethod
    async def update_certification(session: AsyncSession, user_id: UUID, certification_id: str, data: CandidateCertificationUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        updates = {k: v for k, v in data.model_dump().items() if k in data.model_fields_set}
        updated = await CandidateProfileRepo.update_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="certifications_json",
            list_key="certifications",
            item_id_field="certification_id",
            item_id=certification_id,
            updates=updates,
        )
        if updated is None:
            raise HTTPException(status_code=404, detail="Certification not found")
        return {"message": "Certification updated successfully", "certification": updated}

    @staticmethod
    async def delete_certification(session: AsyncSession, user_id: UUID, certification_id: str) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)
        deleted = await CandidateProfileRepo.delete_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="certifications_json",
            list_key="certifications",
            item_id_field="certification_id",
            item_id=certification_id,
        )
        if not deleted:
            raise HTTPException(status_code=404, detail="Certification not found")
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Certification removed successfully", "profile_completion_pct": completion}

    # ── Projects (Add Project modal: per-entry CRUD) ───────────────────────

    @staticmethod
    async def add_project(session: AsyncSession, user_id: UUID, data: CandidateProjectCreateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)

        item = data.model_dump()
        item["project_id"] = uuid4().hex
        await CandidateProfileRepo.add_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="projects_json",
            list_key="projects",
            item=item,
        )
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Project added successfully", "profile_completion_pct": completion, "project": item}

    @staticmethod
    async def update_project(session: AsyncSession, user_id: UUID, project_id: str, data: CandidateProjectUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        updated = await CandidateProfileRepo.update_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="projects_json",
            list_key="projects",
            item_id_field="project_id",
            item_id=project_id,
            updates=data.model_dump(exclude_unset=True),
        )
        if updated is None:
            raise HTTPException(status_code=404, detail="Project not found")
        return {"message": "Project updated successfully", "project": updated}

    @staticmethod
    async def delete_project(session: AsyncSession, user_id: UUID, project_id: str) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)
        deleted = await CandidateProfileRepo.delete_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="projects_json",
            list_key="projects",
            item_id_field="project_id",
            item_id=project_id,
        )
        if not deleted:
            raise HTTPException(status_code=404, detail="Project not found")
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Project removed successfully", "profile_completion_pct": completion}

    # ── Languages (Add Language modal: per-entry CRUD) ─────────────────────

    @staticmethod
    async def add_language(session: AsyncSession, user_id: UUID, data: CandidateLanguageCreateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)

        item = data.model_dump()
        item["language_id"] = uuid4().hex
        await CandidateProfileRepo.add_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="languages_json",
            list_key="languages",
            item=item,
        )
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Language added successfully", "profile_completion_pct": completion, "language": item}

    @staticmethod
    async def update_language(session: AsyncSession, user_id: UUID, language_id: str, data: CandidateLanguageUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        updated = await CandidateProfileRepo.update_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="languages_json",
            list_key="languages",
            item_id_field="language_id",
            item_id=language_id,
            updates=data.model_dump(exclude_unset=True),
        )
        if updated is None:
            raise HTTPException(status_code=404, detail="Language not found")
        return {"message": "Language updated successfully", "language": updated}

    @staticmethod
    async def delete_language(session: AsyncSession, user_id: UUID, language_id: str) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        user = await _get_user_or_404(session, user_id)
        deleted = await CandidateProfileRepo.delete_json_list_item(
            session=session,
            candidate_id=profile.candidate_id,
            json_column="languages_json",
            list_key="languages",
            item_id_field="language_id",
            item_id=language_id,
        )
        if not deleted:
            raise HTTPException(status_code=404, detail="Language not found")
        resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
        completion = _calc_completion(profile, user, resume_detail)
        await CandidateProfileRepo.update_profile(
            session=session, candidate_id=profile.candidate_id, updated_by=str(user_id),
            profile_completion_pct=completion,
        )
        return {"message": "Language removed successfully", "profile_completion_pct": completion}


    @staticmethod
    async def update_visibility(session: AsyncSession, user_id: UUID, data: CandidateVisibilityUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)

        update_kwargs = {k: v for k, v in data.model_dump().items() if v is not None}
        if not update_kwargs:
            return {"message": "No changes provided"}

        requires_visibility_feature = (
            update_kwargs.get("open_to_work") is True
            or update_kwargs.get("searchable_flag") is True
            or update_kwargs.get("profile_visibility") in {"PUBLIC", "CONNECTIONS"}
        )
        if requires_visibility_feature:
            await SubscriptionValidator(
                session=session,
                user_id=user_id,
                role="CANDIDATE",
            ).require_feature("resume_visibility")

        await CandidateProfileRepo.update_profile(
            session=session,
            candidate_id=profile.candidate_id,
            updated_by=str(user_id),
            **update_kwargs,
        )
        return {
            "message": "Visibility settings updated successfully",
            "profile_visibility": update_kwargs.get("profile_visibility", profile.profile_visibility),
            "searchable_flag": update_kwargs.get("searchable_flag", profile.searchable_flag),
            "open_to_work": update_kwargs.get("open_to_work", profile.open_to_work),
            "search_engine_indexing": update_kwargs.get(
                "search_engine_indexing",
                getattr(profile, "search_engine_indexing", False),
            ),
        }

    @staticmethod
    async def update_professional_snapshot(session: AsyncSession, user_id: UUID, data: CandidateProfessionalSnapshotUpdateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        sent = data.model_fields_set

        update_kwargs = {}
        for field in (
            "experience_level", "current_company", "desired_employment",
            "work_preference", "current_ctc", "expected_ctc",
            "notice_period", "salary_expectation", "target_roles",
        ):
            if field in sent:
                update_kwargs[field] = getattr(data, field)

        if "notice_period_id" in sent:
            if data.notice_period_id is not None:
                notice_period = await MasterDataService.validate_notice_period(session, data.notice_period_id)
                update_kwargs["notice_period_id"] = notice_period.notice_period_id
                update_kwargs["notice_period"] = notice_period.label  # keep display-name cache in sync
            else:
                # Explicit clear -- null out both the id and the cached
                # display-name text, otherwise the stale id survives and
                # get_full_profile resolves it right back to the old label.
                update_kwargs["notice_period_id"] = None
                update_kwargs["notice_period"] = None

        if "salary_expectation_id" in sent:
            if data.salary_expectation_id is not None:
                salary_expectation = await MasterDataService.validate_salary_expectation(session, data.salary_expectation_id)
                update_kwargs["salary_expectation_id"] = salary_expectation.salary_expectation_id
                update_kwargs["salary_expectation"] = salary_expectation.label  # keep display-name cache in sync
            else:
                update_kwargs["salary_expectation_id"] = None
                update_kwargs["salary_expectation"] = None
        target_roles_provided = ("target_role_ids" in sent) or ("target_role_names" in sent)
        if target_roles_provided:
            resolved_ids = await MasterDataService.resolve_target_roles(
                session, data.target_role_ids, data.target_role_names
            )
            await MasterDataRepo.replace_candidate_target_roles(session, profile.candidate_id, resolved_ids)

        if not update_kwargs and not target_roles_provided:
            return {"message": "No changes provided"}

        if update_kwargs:
            await CandidateProfileRepo.update_profile(
                session=session,
                candidate_id=profile.candidate_id,
                updated_by=str(user_id),
                **update_kwargs,
            )
        return {"message": "Professional snapshot updated successfully"}

    @staticmethod
    async def upload_resume_file(
        session: AsyncSession,
        user_id: UUID,
        filename: str | None,
        content_type: str | None,
        contents: bytes,
        check_duplicate: bool = True,
    ) -> CandidateResumeResponse:
        extension = _safe_resume_extension(filename, content_type)
        _validate_resume_filename_length(filename)
        if len(contents) < MIN_RESUME_SIZE_BYTES:
            raise HTTPException(status_code=400, detail="The selected file is empty. Please upload a valid resume file.")
        if len(contents) > MAX_RESUME_SIZE_BYTES:
            raise HTTPException(status_code=400, detail="File size must not exceed 5 MB")
        _validate_resume_has_content(contents, extension)

        profile = await _get_profile_or_404(session, user_id)

        existing_resumes = await CandidateProfileRepo.get_resumes(session, profile.candidate_id)

        file_hash = hashlib.sha256(contents).hexdigest()
        clean_name = Path(filename or f"resume{extension}").name
        if check_duplicate:
            if any(existing.file_hash == file_hash for existing in existing_resumes):
                raise HTTPException(
                    status_code=400,
                    detail="You've already uploaded this file. Please choose a different file.",
                )
            if any(existing.file_name.lower() == clean_name.lower() for existing in existing_resumes):
                raise HTTPException(
                    status_code=400,
                    detail="A resume with this file name already exists. Please rename the file and try again.",
                )

        validator = SubscriptionValidator(
            session=session,
            user_id=user_id,
            role="CANDIDATE",
        )
        subscription_resume_limit = await validator.get_limit("max_resume_uploads")
        # Every candidate is guaranteed at least MAX_RESUME_COUNT (5) resumes
        # regardless of subscription plan configuration -- a plan can only
        # raise this ceiling (e.g. a premium tier offering more), never lower
        # it below the platform floor. Without this floor, a plan record
        # with a smaller/misconfigured max_resume_uploads (or one that
        # predates a rollout of a higher default) would incorrectly cap a
        # candidate below the number the "Attached CV" panel already
        # advertises (see MAX_RESUME_COUNT in EditProfileResumeBuilder.jsx).
        effective_resume_limit = max(
            subscription_resume_limit if subscription_resume_limit is not None else MAX_RESUME_COUNT,
            MAX_RESUME_COUNT,
        )

        if len(existing_resumes) >= effective_resume_limit:
            raise HTTPException(
                status_code=403,
                detail=f"You can only upload up to {effective_resume_limit} resumes. Please delete an existing resume before uploading another.",
            )

        await session.execute(
            text("UPDATE candidate_resumes SET is_active = FALSE WHERE candidate_id = :cid"),
            {"cid": profile.candidate_id},
        )
        await session.commit()

        new_resume = CandidateResume(
            candidate_id=profile.candidate_id,
            file_name=clean_name,
            file_size=len(contents),
            file_hash=file_hash,
            version_name=Path(clean_name).stem,
            is_active=True,
            created_by=str(user_id),
        )
        session.add(new_resume)
        await session.flush()

        s3_key = s3_service.upload_resume(
            contents=contents,
            user_id=str(user_id),
            resume_id=new_resume.resume_id,
            extension=extension,
            content_type=content_type or "application/octet-stream",
        )

        new_resume.file_path = None          
        new_resume.blob_ref = s3_key
        await session.commit()

        await CandidateProfileRepo.update_profile(
            session=session,
            candidate_id=profile.candidate_id,
            updated_by=str(user_id),
            active_resume_id=new_resume.resume_id,
        )
        await ActivityLogService.create_log_for_user_id(
            session=session,
            user_id=user_id,
            actor_role="CANDIDATE",
            action="RESUME_UPLOADED",
            entity_type="Resume",
            entity_id=new_resume.resume_id,
            target_entity_name=new_resume.file_name,
            description=f"Uploaded resume {new_resume.file_name}",
            metadata={"candidate_id": profile.candidate_id},
        )

        return CandidateResumeResponse.model_validate(new_resume)

    @staticmethod
    async def _resolve_default_resume(
        session: AsyncSession, profile, user_id: UUID, resumes: Optional[list[CandidateResume]] = None
    ) -> tuple[Optional[str], bool]:
        if profile.active_resume_id:
            current = await CandidateProfileRepo.get_resume(session, profile.candidate_id, profile.active_resume_id)
            if current and not current.is_archived:
                return profile.active_resume_id, True

        fallback = await CandidateProfileRepo.get_most_recent_resume(session, profile.candidate_id)
        if not fallback:
            return None, False

        if fallback.resume_id != profile.active_resume_id:
            await CandidateProfileRepo.set_active_resume(
                session=session, candidate_id=profile.candidate_id, resume_id=fallback.resume_id, user_id=user_id
            )
            profile.active_resume_id = fallback.resume_id

        return fallback.resume_id, False

    @staticmethod
    async def list_resumes(session: AsyncSession, user_id: UUID) -> CandidateResumeListResponse:
        profile = await _get_profile_or_404(session, user_id)
        resumes = await CandidateProfileRepo.get_resumes(session, profile.candidate_id)
        default_resume_id, is_explicit = await CandidateProfileService._resolve_default_resume(
            session, profile, user_id, resumes
        )
        if default_resume_id and not is_explicit:
            for resume in resumes:
                resume.is_active = resume.resume_id == default_resume_id
        return CandidateResumeListResponse(
            total=len(resumes),
            active_resume_id=default_resume_id,
            default_is_explicit=is_explicit,
            items=[CandidateResumeResponse.model_validate(resume) for resume in resumes],
        )

    @staticmethod
    async def set_active_resume(session: AsyncSession, user_id: UUID, resume_id: str) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        resume = await CandidateProfileRepo.get_resume(session, profile.candidate_id, resume_id)
        if not resume:
            raise HTTPException(status_code=404, detail="Resume not found")
        if resume.is_archived:
            raise HTTPException(status_code=400, detail="Archived resumes can't be set as default. Restore it first.")
        updated = await CandidateProfileRepo.set_active_resume(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
        )
        if not updated:
            raise HTTPException(status_code=404, detail="Resume not found")
        return {"message": "Active resume updated successfully", "resume_id": resume_id}

    @staticmethod
    async def delete_resume(session: AsyncSession, user_id: UUID, resume_id: str) -> dict:
        profile = await _get_profile_or_404(session, user_id)
        resume = await CandidateProfileRepo.get_resume(session, profile.candidate_id, resume_id)
        deleted = await CandidateProfileRepo.delete_resume(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
        )
        if not deleted:
            raise HTTPException(status_code=404, detail="Resume not found")
        await ActivityLogService.create_log_for_user_id(
            session=session,
            user_id=user_id,
            actor_role="CANDIDATE",
            action="RESUME_DELETED",
            entity_type="Resume",
            entity_id=resume_id,
            target_entity_name=getattr(resume, "file_name", None),
            description=f"Deleted resume {getattr(resume, 'file_name', resume_id)}",
            metadata={"candidate_id": profile.candidate_id},
        )
        return {"message": "Resume deleted successfully"}

    @staticmethod
    async def get_resume_download(
        session: AsyncSession,
        user_id: UUID,
        resume_id: Optional[str] = None,
        count_as_download: bool = True,
        as_attachment: bool = True,
    ) -> CandidateResumeDownloadResponse:
        profile = await _get_profile_or_404(session, user_id)
        target_resume_id = resume_id
        if not target_resume_id:
            target_resume_id, _ = await CandidateProfileService._resolve_default_resume(session, profile, user_id)
        if not target_resume_id:
            raise HTTPException(
                status_code=404,
                detail="No default resume available. Please upload a resume.",
            )
        resume = await CandidateProfileRepo.get_resume(session, profile.candidate_id, target_resume_id)
        if not resume:
            raise HTTPException(status_code=404, detail="Resume not found")
        if resume.blob_ref:
            presigned = s3_service.generate_presigned_url(
                resume.blob_ref,
                expiry_seconds=3600,
                attachment_filename=resume.file_name if as_attachment else None,
            )
            resume.file_path = presigned
        if count_as_download:
            await CandidateProfileRepo.increment_resume_download_count(session, target_resume_id)
        return CandidateResumeDownloadResponse.model_validate(resume)

    @staticmethod
    async def get_resume_download_file(
        session: AsyncSession,
        user_id: UUID,
        resume_id: Optional[str] = None,
        count_as_download: bool = True,
    ) -> tuple[bytes, str, str | None]:
        
        profile = await _get_profile_or_404(session, user_id)
        target_resume_id = resume_id
        if not target_resume_id:
            target_resume_id, _ = await CandidateProfileService._resolve_default_resume(session, profile, user_id)
        if not target_resume_id:
            raise HTTPException(
                status_code=404,
                detail="No default resume available. Please upload a resume.",
            )
        resume = await CandidateProfileRepo.get_resume(session, profile.candidate_id, target_resume_id)
        if not resume:
            raise HTTPException(status_code=404, detail="Resume not found")
        if not resume.blob_ref:
            raise HTTPException(status_code=404, detail="Resume file not found")

        content, content_type = s3_service.fetch_object_bytes(resume.blob_ref)
        if count_as_download:
            await CandidateProfileRepo.increment_resume_download_count(session, target_resume_id)
        return content, (resume.file_name or "resume"), content_type

    @staticmethod
    async def parse_resume_content(
        session: AsyncSession, user_id: UUID, resume_id: str
    ) -> dict:
        
        content, file_name, content_type = await CandidateProfileService.get_resume_download_file(
            session, user_id, resume_id, count_as_download=False,
        )
        extracted = resume_extraction.extract_resume_data(content, file_name, content_type)
        if not extracted:
            raise HTTPException(
                status_code=422,
                detail="Could not extract data from this resume file. It may be an image-only scan or an unsupported format.",
            )
        return extracted

    @staticmethod
    async def get_resume_preview(
        session: AsyncSession, user_id: UUID, resume_id: Optional[str] = None
    ) -> CandidateResumeDownloadResponse:
        
        return await CandidateProfileService.get_resume_download(
            session, user_id, resume_id, count_as_download=False, as_attachment=False
        )

    @staticmethod
    async def update_resume(
        session: AsyncSession, user_id: UUID, resume_id: str, payload: CandidateResumeUpdateSchema
    ) -> CandidateResumeResponse:
        profile = await _get_profile_or_404(session, user_id)
        updated = await CandidateProfileRepo.update_resume_meta(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
            values=payload.model_dump(exclude_unset=True),
        )
        if not updated:
            raise HTTPException(status_code=404, detail="Resume not found")
        return CandidateResumeResponse.model_validate(updated)

    @staticmethod
    async def set_resume_archived(session: AsyncSession, user_id: UUID, resume_id: str, archived: bool) -> CandidateResumeResponse:
        profile = await _get_profile_or_404(session, user_id)
        updated = await CandidateProfileRepo.set_resume_archived(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
            archived=archived,
        )
        if not updated:
            raise HTTPException(status_code=404, detail="Resume not found")
        return CandidateResumeResponse.model_validate(updated)

    @staticmethod
    async def duplicate_resume(session: AsyncSession, user_id: UUID, resume_id: str) -> CandidateResumeResponse:
        profile = await _get_profile_or_404(session, user_id)
        existing_resumes = await CandidateProfileRepo.get_resumes(session, profile.candidate_id)
        validator = SubscriptionValidator(
            session=session,
            user_id=user_id,
            role="CANDIDATE",
        )
        subscription_resume_limit = await validator.get_limit("max_resume_uploads")
        # Same platform floor as upload_resume_file above -- a subscription
        # plan can raise this above MAX_RESUME_COUNT, never lower it below.
        effective_resume_limit = max(
            subscription_resume_limit if subscription_resume_limit is not None else MAX_RESUME_COUNT,
            MAX_RESUME_COUNT,
        )
        if len(existing_resumes) >= effective_resume_limit:
            raise HTTPException(
                status_code=403,
                detail=f"You can only have up to {effective_resume_limit} resumes. Please delete an existing resume before duplicating.",
            )
        copy = await CandidateProfileRepo.duplicate_resume(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
        )
        if not copy:
            raise HTTPException(status_code=404, detail="Resume not found")
        return CandidateResumeResponse.model_validate(copy)

    @staticmethod
    async def create_resume_share_link(
        session: AsyncSession, user_id: UUID, resume_id: str, payload: ResumeShareLinkCreateSchema
    ) -> ResumeShareLinkResponse:
        from datetime import timedelta

        profile = await _get_profile_or_404(session, user_id)
        share_token = uuid4().hex
        expires_at = utc_now_naive() + timedelta(days=payload.expires_in_days) if payload.expires_in_days else None
        updated = await CandidateProfileRepo.set_resume_share_link(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
            share_token=share_token,
            requires_email=payload.requires_email,
            expires_at=expires_at,
        )
        if not updated:
            raise HTTPException(status_code=404, detail="Resume not found")
        return _resume_share_response(updated)

    @staticmethod
    async def disable_resume_share_link(session: AsyncSession, user_id: UUID, resume_id: str) -> ResumeShareLinkResponse:
        profile = await _get_profile_or_404(session, user_id)
        updated = await CandidateProfileRepo.disable_resume_share_link(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
        )
        if not updated:
            raise HTTPException(status_code=404, detail="Resume not found")
        return _resume_share_response(updated)

    @staticmethod
    async def enable_resume_share_link(session: AsyncSession, user_id: UUID, resume_id: str) -> ResumeShareLinkResponse:
        profile = await _get_profile_or_404(session, user_id)
        resume = await CandidateProfileRepo.get_resume(session, profile.candidate_id, resume_id)
        if not resume:
            raise HTTPException(status_code=404, detail="Resume not found")
        if not resume.share_token:
            raise HTTPException(status_code=400, detail="No share link to enable. Create one first.")
        if resume.share_expires_at and resume.share_expires_at <= utc_now_naive():
            raise HTTPException(status_code=400, detail="This link has expired. Create a new one instead.")

        updated = await CandidateProfileRepo.enable_resume_share_link(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
        )
        if not updated:
            raise HTTPException(status_code=404, detail="Resume not found")
        return _resume_share_response(updated)

    @staticmethod
    async def delete_resume_share_link(session: AsyncSession, user_id: UUID, resume_id: str) -> ResumeShareLinkResponse:
        profile = await _get_profile_or_404(session, user_id)
        updated = await CandidateProfileRepo.delete_resume_share_link(
            session=session,
            candidate_id=profile.candidate_id,
            resume_id=resume_id,
            user_id=user_id,
        )
        if not updated:
            raise HTTPException(status_code=404, detail="Resume not found")
        return _resume_share_response(updated)

    # -----------------------------------------------------------------------
    # Public share link resolution (unauthenticated /cv/{token} page)
    # -----------------------------------------------------------------------

    @staticmethod
    async def _resolve_shared_resume(session: AsyncSession, share_token: str) -> CandidateResume:
        resume = await CandidateProfileRepo.get_resume_by_share_token(session, share_token)
        if not resume or not resume.share_enabled or resume.is_archived:
            raise HTTPException(status_code=404, detail="This link is invalid or no longer available")
        if resume.share_expires_at and resume.share_expires_at <= utc_now_naive():
            raise HTTPException(status_code=404, detail="This link has expired")
        return resume

    @staticmethod
    async def get_shared_resume_info(session: AsyncSession, share_token: str) -> SharedResumeInfoResponse:
        resume = await CandidateProfileService._resolve_shared_resume(session, share_token)

        name_row = await session.execute(
            text(
                """
                SELECT u.first_name, u.last_name
                FROM candidate_profiles cp
                JOIN users u ON u.user_id = cp.user_id
                WHERE cp.candidate_id = :cid
                """
            ),
            {"cid": resume.candidate_id},
        )
        name = name_row.first()
        candidate_name = " ".join(part for part in [name.first_name if name else None, name.last_name if name else None] if part) or "A candidate"

        download_url = None
        if not resume.share_requires_email:
            if resume.blob_ref:
                download_url = s3_service.generate_presigned_url(resume.blob_ref, expiry_seconds=3600)
            await CandidateProfileRepo.increment_resume_share_view_count(session, resume.resume_id)

        return SharedResumeInfoResponse(
            candidate_name=candidate_name,
            file_name=resume.file_name,
            template=resume.template,
            requires_email=resume.share_requires_email,
            download_url=download_url,
        )

    @staticmethod
    async def access_shared_resume(session: AsyncSession, share_token: str, email: str) -> SharedResumeAccessResponse:
        resume = await CandidateProfileService._resolve_shared_resume(session, share_token)
        if not resume.blob_ref:
            raise HTTPException(status_code=404, detail="This resume file is no longer available")

        download_url = s3_service.generate_presigned_url(resume.blob_ref, expiry_seconds=3600)
        await CandidateProfileRepo.increment_resume_share_view_count(session, resume.resume_id)

        return SharedResumeAccessResponse(download_url=download_url, file_name=resume.file_name)

    @staticmethod
    async def generate_resume_pdf(
        session: AsyncSession, user_id: UUID, style: str = "ats", resume_id: Optional[str] = None
    ) -> tuple[bytes, str]:
        await SubscriptionValidator(
            session=session,
            user_id=user_id,
            role="CANDIDATE",
        ).require_feature("resume_builder")

        profile = await _get_profile_or_404(session, user_id)
        user = await CandidateProfileRepo.get_user_by_id(session, user_id)

        resume = None
        extracted = None
        if resume_id:
            resume = await CandidateProfileRepo.get_resume(session, profile.candidate_id, resume_id)
            if resume and resume.blob_ref:
                try:
                    content, content_type = s3_service.fetch_object_bytes(resume.blob_ref)
                    extracted = resume_extraction.extract_resume_data(
                        content, resume.file_name or "", content_type
                    )
                except HTTPException:
                    extracted = None

        if extracted:
            pdf_bytes = resume_document_service.render_resume_pdf_from_extracted(
                extracted, style=style, fallback_user=user
            )
        else:
            resume_detail = await CandidateProfileRepo.get_latest_resume_detail(session, profile.candidate_id)
            pdf_bytes = resume_document_service.render_resume_pdf(profile, user, resume_detail, style=style)

        
        filename = f"resume-{style}.pdf"
        if resume_id and resume:
            await CandidateProfileRepo.increment_resume_download_count(session, resume_id)
        return pdf_bytes, filename

    @staticmethod
    async def generate_resume_studio_pdf(
        session: AsyncSession, user_id: UUID, draft: dict, style: str = "ats"
    ) -> tuple[bytes, str]:
        await SubscriptionValidator(
            session=session,
            user_id=user_id,
            role="CANDIDATE",
        ).require_feature("resume_builder")

        user = await CandidateProfileRepo.get_user_by_id(session, user_id)
        extracted = resume_document_service.draft_to_extracted(draft)
        pdf_bytes = resume_document_service.render_resume_pdf_from_extracted(
            extracted, style=style, fallback_user=user
        )
        fallback_name = f"{(user.first_name or '').strip()} {(user.last_name or '').strip()}".strip()
        name = extracted["full_name"] or fallback_name or "resume"
        filename = f"{name.lower().replace(' ', '-')}-{style}.pdf"
        return pdf_bytes, filename

    @staticmethod
    async def list_saved_jobs(session: AsyncSession, user_id: UUID, page: int, page_size: int) -> CandidateSavedJobListResponse:
        profile = await _get_profile_or_404(session, user_id)
        total, rows, skills_by_job = await CandidateProfileRepo.get_saved_jobs(
            session=session,
            candidate_id=profile.candidate_id,
            page=page,
            page_size=page_size,
        )
        return CandidateSavedJobListResponse(
            total=total,
            page=page,
            page_size=page_size,
            items=[
                CandidateSavedJobResponse(
                    saved_job_id=saved.saved_job_id,
                    job_id=saved.job_id,
                    saved_flag=saved.saved_flag,
                    created_at=saved.created_at,
                    job_title=job.title,
                    company_name=job.company_name,
                    location=job.location,
                    employment_type=job.employment_type,
                    work_mode=job.work_mode,
                    # Feeds the "stacked by role" grouping on the Favourites
                    # page (e.g. Java / Python), replacing the old hardcoded
                    # Collections list, instead of a made-up category.
                    skills=skills_by_job.get(job.job_id, []),
                )
                for saved, job in rows
            ],
        )

    @staticmethod
    async def save_job(session: AsyncSession, user_id: UUID, data: CandidateSavedJobCreateSchema) -> dict:
        profile = await _get_profile_or_404(session, user_id)

        MAX_SAVED_JOBS = 30
        validator = SubscriptionValidator(
            session=session,
            user_id=user_id,
            role="CANDIDATE",
        )
        await validator.require_feature("saved_jobs")
        already_saved = await CandidateProfileRepo.is_job_saved(session, profile.candidate_id, data.job_id)
        if not already_saved:
            current_count = await CandidateProfileRepo.count_saved_jobs(session, profile.candidate_id)
            subscription_saved_job_limit = await validator.get_limit("saved_jobs_limit")
            effective_saved_job_limit = (
                subscription_saved_job_limit
                if subscription_saved_job_limit is not None
                else MAX_SAVED_JOBS
            )
            if current_count >= effective_saved_job_limit:
                raise HTTPException(
                    status_code=403,
                    detail=f"You can save up to {effective_saved_job_limit} jobs. Remove a saved job before adding a new one.",
                )

        saved = await CandidateProfileRepo.save_job(session, profile.candidate_id, data.job_id)
        return {"message": "Job saved successfully", "saved_job_id": saved.saved_job_id, "job_id": saved.job_id}

    @staticmethod
    async def unsave_job(session: AsyncSession, user_id: UUID, job_id: str) -> dict:
        await _require_candidate_feature(session, user_id, "saved_jobs")
        profile = await _get_profile_or_404(session, user_id)
        deleted = await CandidateProfileRepo.unsave_job(session, profile.candidate_id, job_id)
        if not deleted:
            raise HTTPException(status_code=404, detail="Saved job not found")
        return {"message": "Job removed from saved jobs"}

    @staticmethod
    async def list_job_alerts(session: AsyncSession, user_id: UUID) -> list[JobAlertResponse]:
        profile = await _get_profile_or_404(session, user_id)
        alerts = await CandidateProfileRepo.get_job_alerts(session, profile.candidate_id)
        return [JobAlertResponse.model_validate(alert) for alert in alerts]

    @staticmethod
    async def create_job_alert(
        session: AsyncSession,
        user_id: UUID,
        data: JobAlertUpsertSchema,
    ) -> JobAlertResponse:

        profile = await _get_profile_or_404(session, user_id)
        validator = SubscriptionValidator(
            session=session,
            user_id=user_id,
            role="CANDIDATE",
        )
        await validator.require_feature("job_alerts")
        max_job_alerts = await validator.get_limit("max_job_alerts")
        if max_job_alerts is not None:
            current_alert_count = len(
                await CandidateProfileRepo.get_job_alerts(session, profile.candidate_id)
            )
            if current_alert_count >= max_job_alerts:
                raise HTTPException(
                    status_code=403,
                    detail=f"You can create up to {max_job_alerts} job alerts. Delete an existing alert before creating another.",
                )
        if data.job_category:
            await MasterDataService.validate_job_category(
                session=session,
                job_category=data.job_category,
            )

        existing_title = await CandidateProfileRepo.find_job_alert_by_title(
            session=session,
            candidate_id=profile.candidate_id,
            title=data.title,
        )
        if existing_title:
            raise HTTPException(
                status_code=409,
                detail=JOB_ALERT_TITLE_DUPLICATE_MESSAGE,
            )

        existing = await CandidateProfileRepo.find_duplicate_job_alert(
            session=session,
            candidate_id=profile.candidate_id,
            title=data.title,
            job_category=data.job_category,
            job_title=data.job_title,
            preferred_location=data.preferred_location,
            experience_level=data.experience_level,
            employment_type=data.employment_type,
        )

        if existing:
            raise HTTPException(
                status_code=409,
                detail="Job alert already exists.",
            )

        create_data = data.model_dump()
        create_data["timezone"] = create_data.get("timezone") or JOB_ALERT_DEFAULT_TIMEZONE

        alert = await CandidateProfileRepo.create_job_alert(
            session=session,
            candidate_id=profile.candidate_id,
            **create_data,
        )
        await ActivityLogService.create_log_for_user_id(
            session=session,
            user_id=user_id,
            actor_role="CANDIDATE",
            action="JOB_ALERT_CREATED",
            entity_type="JobAlert",
            entity_id=alert.alert_id,
            target_entity_name=alert.title,
            description=f"Created job alert {alert.title}",
        )

        return JobAlertResponse.model_validate(alert)

    @staticmethod
    async def update_job_alert(
        session: AsyncSession,
        user_id: UUID,
        alert_id: str,
        data: JobAlertUpdateSchema,
    ) -> JobAlertResponse:

        await _require_candidate_feature(session, user_id, "job_alerts")
        profile = await _get_profile_or_404(session, user_id)

        # Fetch existing alert
        existing_alert = await CandidateProfileRepo.get_job_alert(
            session=session,
            candidate_id=profile.candidate_id,
            alert_id=alert_id,
        )

        if not existing_alert:
            raise HTTPException(
                status_code=404,
                detail="Job alert not found",
            )

        # Only include fields supplied in request
        update_data = data.model_dump(
            exclude_unset=True,
        )
        if "timezone" not in update_data and not getattr(existing_alert, "timezone", None):
            update_data["timezone"] = JOB_ALERT_DEFAULT_TIMEZONE
        elif "timezone" in update_data and update_data["timezone"] is None:
            update_data["timezone"] = JOB_ALERT_DEFAULT_TIMEZONE

        # Merge existing values with incoming values
        merged_title = update_data.get(
            "title",
            existing_alert.title,
        )

        merged_job_category = update_data.get(
            "job_category",
            existing_alert.job_category,
        )

        merged_job_title = update_data.get(
            "job_title",
            existing_alert.job_title,
        )

        merged_preferred_location = update_data.get(
            "preferred_location",
            existing_alert.preferred_location,
        )

        merged_experience_level = update_data.get(
            "experience_level",
            existing_alert.experience_level,
        )

        merged_employment_type = update_data.get(
            "employment_type",
            existing_alert.employment_type,
        )

        if merged_job_category:
            await MasterDataService.validate_job_category(
                session=session,
                job_category=merged_job_category,
            )

        existing_title = await CandidateProfileRepo.find_job_alert_by_title(
            session=session,
            candidate_id=profile.candidate_id,
            title=merged_title,
            exclude_alert_id=alert_id,
        )
        if existing_title:
            raise HTTPException(
                status_code=409,
                detail=JOB_ALERT_TITLE_DUPLICATE_MESSAGE,
            )

        # Duplicate check
        duplicate = await CandidateProfileRepo.find_duplicate_job_alert(
            session=session,
            candidate_id=profile.candidate_id,
            title=merged_title,
            job_category=merged_job_category,
            job_title=merged_job_title,
            preferred_location=merged_preferred_location,
            experience_level=merged_experience_level,
            employment_type=merged_employment_type,
        )

        if duplicate and duplicate.alert_id != alert_id:
            raise HTTPException(
                status_code=409,
                detail="Job alert already exists.",
            )

        # Update only supplied fields
        alert = await CandidateProfileRepo.update_job_alert(
            session=session,
            candidate_id=profile.candidate_id,
            alert_id=alert_id,
            **update_data,
        )

        if not alert:
            raise HTTPException(
                status_code=404,
                detail="Job alert not found",
            )

        await ActivityLogService.create_log_for_user_id(
            session=session,
            user_id=user_id,
            actor_role="CANDIDATE",
            action="JOB_ALERT_UPDATED",
            entity_type="JobAlert",
            entity_id=alert.alert_id,
            target_entity_name=alert.title,
            description=f"Updated job alert {alert.title}",
        )
        return JobAlertResponse.model_validate(alert)

    @staticmethod
    async def delete_job_alert(session: AsyncSession, user_id: UUID, alert_id: str) -> dict:
        if not JOB_ALERT_ID_PATTERN.fullmatch(alert_id or ""):
            raise HTTPException(status_code=400, detail="Invalid job alert ID")

        await _require_candidate_feature(session, user_id, "job_alerts")
        profile = await _get_profile_or_404(session, user_id)
        deleted = await CandidateProfileRepo.delete_job_alert(session, profile.candidate_id, alert_id)
        if not deleted:
            raise HTTPException(status_code=404, detail="Job alert not found")
        await ActivityLogService.create_log_for_user_id(
            session=session,
            user_id=user_id,
            actor_role="CANDIDATE",
            action="JOB_ALERT_DELETED",
            entity_type="JobAlert",
            entity_id=alert_id,
            description=f"Deleted job alert {alert_id}",
        )
        return {"message": "Job alert deleted successfully"}

    @staticmethod
    async def list_saved_searches(session: AsyncSession, user_id: UUID) -> list[SavedSearchResponse]:
        profile = await _get_profile_or_404(session, user_id)
        searches = await CandidateProfileRepo.get_saved_searches(session, profile.candidate_id)
        return [SavedSearchResponse.model_validate(search) for search in searches]

    @staticmethod
    async def create_saved_search(session: AsyncSession, user_id: UUID, data: SavedSearchCreateSchema) -> SavedSearchResponse:
        profile = await _get_profile_or_404(session, user_id)
        saved_search = await CandidateProfileRepo.create_saved_search(
            session=session,
            candidate_id=profile.candidate_id,
            **data.model_dump(),
        )
        return SavedSearchResponse.model_validate(saved_search)

    # -----------------------------------------------------------------------
    # Profile picture & cover picture upload (S3)
    # -----------------------------------------------------------------------

    ALLOWED_IMAGE_TYPES: dict = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/webp": ".webp",
    }
    MAX_IMAGE_SIZE_BYTES: int = 5 * 1024 * 1024  # 5 MB

    @staticmethod
    async def upload_profile_picture(
        session: AsyncSession,
        user_id: UUID,
        filename: str | None,
        content_type: str | None,
        contents: bytes,
    ) -> dict:
        extension = CandidateProfileService.ALLOWED_IMAGE_TYPES.get(content_type or "")
        if not extension:
            raise HTTPException(status_code=400, detail="Only JPEG, PNG, and WebP images are allowed")
        if len(contents) > CandidateProfileService.MAX_IMAGE_SIZE_BYTES:
            raise HTTPException(status_code=400, detail="Image size must not exceed 5 MB")

        s3_key = s3_service.upload_profile_picture(
            contents=contents,
            user_id=str(user_id),
            extension=extension,
            content_type=content_type,
        )
        presigned_url = s3_service.generate_presigned_url(s3_key, expiry_seconds=3600)

        # Persist the S3 key as the profile_image_url on the users record
        await session.execute(
            text("UPDATE users SET profile_image_url = :url WHERE user_id = :uid"),
            {"url": s3_key, "uid": user_id},
        )
        await session.commit()

        return {"message": "Profile picture uploaded successfully", "url": presigned_url}

    @staticmethod
    async def remove_profile_picture(session: AsyncSession, user_id: UUID) -> dict:
        row = await session.execute(
            text("SELECT profile_image_url FROM users WHERE user_id = :uid"),
            {"uid": user_id},
        )
        current = row.scalar_one_or_none()

        if current:
            
            try:
                s3_service.delete_object(current)
            except Exception:
                pass

        await session.execute(
            text("UPDATE users SET profile_image_url = NULL WHERE user_id = :uid"),
            {"uid": user_id},
        )
        await session.commit()

        return {"message": "Profile picture removed successfully"}

    @staticmethod
    async def upload_cover_picture(
        session: AsyncSession,
        user_id: UUID,
        filename: str | None,
        content_type: str | None,
        contents: bytes,
    ) -> dict:
        extension = CandidateProfileService.ALLOWED_IMAGE_TYPES.get(content_type or "")
        if not extension:
            raise HTTPException(status_code=400, detail="Only JPEG, PNG, and WebP images are allowed")
        if len(contents) > CandidateProfileService.MAX_IMAGE_SIZE_BYTES:
            raise HTTPException(status_code=400, detail="Image size must not exceed 5 MB")

        s3_key = s3_service.upload_cover_picture(
            contents=contents,
            user_id=str(user_id),
            extension=extension,
            content_type=content_type,
        )
        presigned_url = s3_service.generate_presigned_url(s3_key, expiry_seconds=3600)

        # Persist the S3 key as the cover_image_url on the users record
        await session.execute(
            text("UPDATE users SET cover_image_url = :url WHERE user_id = :uid"),
            {"url": s3_key, "uid": user_id},
        )
        await session.commit()

        return {"message": "Cover picture uploaded successfully", "url": presigned_url}

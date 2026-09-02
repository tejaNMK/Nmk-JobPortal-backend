from __future__ import annotations

import logging
import re
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException

from app.schema.job import MAX_JOB_DESCRIPTION_LENGTH

logger = logging.getLogger(__name__)


SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc"}
MAX_DOCUMENT_SIZE_BYTES = 5 * 1024 * 1024
MAX_DOCUMENT_FILENAME_LENGTH = 120
ALLOWED_DOCUMENT_CONTENT_TYPES = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    },
    ".doc": {"application/msword", "application/octet-stream"},
}

_UNREADABLE_DETAIL = "The uploaded document could not be read. Please upload a valid PDF or Word document."


class DocumentTextExtractionService:
    @staticmethod
    def extract_text_from_upload(
        *,
        filename: str | None,
        content_type: str | None,
        contents: bytes,
    ) -> str:
        extension = DocumentTextExtractionService.validate_upload(
            filename=filename,
            contents=contents,
            content_type=content_type,
        )

        try:
            if extension == ".pdf":
                raw_text = DocumentTextExtractionService._extract_pdf_text(contents)
            elif extension == ".docx":
                raw_text = DocumentTextExtractionService._extract_docx_text(contents)
            else:
                raw_text = DocumentTextExtractionService._extract_doc_text(contents)
        except HTTPException:
            raise
        except Exception as exc:
            logger.exception(
                "Failed to extract document text.",
                extra={"filename": filename, "content_type": content_type},
            )
            raise HTTPException(status_code=400, detail=_UNREADABLE_DETAIL) from exc

        text = DocumentTextExtractionService.clean_text(raw_text)
        if not text:
            raise HTTPException(
                status_code=400,
                detail="The uploaded document does not contain readable text.",
            )
        if len(text) > MAX_JOB_DESCRIPTION_LENGTH:
            raise HTTPException(
                status_code=400,
                detail=f"Extracted job description must not exceed {MAX_JOB_DESCRIPTION_LENGTH} characters",
            )
        return text

    @staticmethod
    def validate_upload(
        *,
        filename: str | None,
        contents: bytes,
        content_type: str | None = None,
    ) -> str:
        safe_filename = Path(filename or "").name
        if not safe_filename:
            raise HTTPException(status_code=400, detail="Document file name is required")
        if len(safe_filename) > MAX_DOCUMENT_FILENAME_LENGTH:
            raise HTTPException(
                status_code=400,
                detail=f"Document file name must not exceed {MAX_DOCUMENT_FILENAME_LENGTH} characters",
            )

        extension = Path(safe_filename).suffix.lower()
        if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail="Only PDF, DOCX, and DOC documents are allowed",
            )
        normalized_content_type = (content_type or "").split(";", 1)[0].strip().lower()
        if normalized_content_type:
            allowed_content_types = ALLOWED_DOCUMENT_CONTENT_TYPES[extension]
            if normalized_content_type not in allowed_content_types:
                raise HTTPException(
                    status_code=400,
                    detail="Document content type does not match the selected file type",
                )
        if not contents:
            raise HTTPException(
                status_code=400,
                detail="The selected document is empty. Please upload a valid file.",
            )
        if len(contents) > MAX_DOCUMENT_SIZE_BYTES:
            raise HTTPException(status_code=400, detail="Document size must not exceed 5 MB")
        DocumentTextExtractionService._validate_magic_bytes(extension, contents)
        return extension

    @staticmethod
    def _validate_magic_bytes(extension: str, contents: bytes) -> None:
        if extension == ".pdf" and not contents.startswith(b"%PDF"):
            raise HTTPException(status_code=400, detail=_UNREADABLE_DETAIL)
        if extension == ".docx" and not contents.startswith(b"PK"):
            raise HTTPException(status_code=400, detail=_UNREADABLE_DETAIL)
        if extension == ".doc" and not contents.startswith(b"\xd0\xcf\x11\xe0"):
            raise HTTPException(status_code=400, detail=_UNREADABLE_DETAIL)

    @staticmethod
    def clean_text(text: str | None) -> str:
        if not text:
            return ""

        normalized = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
        lines = []
        for line in normalized.split("\n"):
            cleaned = re.sub(r"[ \t\f\v]+", " ", line).strip()
            lines.append(cleaned)

        paragraphs = []
        current_blank = False
        for line in lines:
            if line:
                paragraphs.append(line)
                current_blank = False
            elif paragraphs and not current_blank:
                paragraphs.append("")
                current_blank = True

        return "\n".join(paragraphs).strip()

    @staticmethod
    def _extract_pdf_text(contents: bytes) -> str:
        try:
            import fitz
        except Exception as exc:
            logger.exception("PyMuPDF is required for PDF document text extraction.")
            raise HTTPException(
                status_code=500,
                detail="PDF document parsing is not configured",
            ) from exc

        try:
            with fitz.open(stream=contents, filetype="pdf") as document:
                return "\n\n".join(page.get_text("text") for page in document)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=_UNREADABLE_DETAIL) from exc

    @staticmethod
    def _extract_docx_text(contents: bytes) -> str:
        try:
            from docx import Document
        except Exception as exc:
            logger.exception("python-docx is required for DOCX document text extraction.")
            raise HTTPException(
                status_code=500,
                detail="DOCX document parsing is not configured",
            ) from exc

        try:
            document = Document(BytesIO(contents))
        except Exception as exc:
            raise HTTPException(status_code=400, detail=_UNREADABLE_DETAIL) from exc

        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    @staticmethod
    def _extract_doc_text(contents: bytes) -> str:
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = Path(tmpdir) / "document.doc"
            doc_path.write_bytes(contents)
            return DocumentTextExtractionService._extract_doc_text_from_path(doc_path)

    @staticmethod
    def _extract_doc_text_from_path(doc_path: Path) -> str:
        try:
            import textract
        except Exception:
            textract = None

        if textract is not None:
            try:
                return textract.process(str(doc_path)).decode("utf-8", errors="ignore")
            except Exception:
                logger.exception("textract failed to parse legacy DOC document.")

        antiword = shutil.which("antiword")
        if antiword:
            try:
                completed = subprocess.run(
                    [antiword, str(doc_path)],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    check=True,
                )
                return completed.stdout.decode("utf-8", errors="ignore")
            except Exception as exc:
                logger.exception("antiword failed to parse legacy DOC document.")
                raise HTTPException(status_code=400, detail=_UNREADABLE_DETAIL) from exc

        raise HTTPException(
            status_code=422,
            detail="Legacy DOC document parsing is not available on this server. Please upload a PDF or DOCX file.",
        )
